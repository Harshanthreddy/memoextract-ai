import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def get_value(field, default="—"):
    """Helper to unwrap field value whether it is a dict {'value': ...} or primitive"""
    if isinstance(field, dict):
        return field.get("value", default)
    if field is None:
        return default
    return field


def create_excel_report(data: dict) -> bytes:
    """
    Takes extracted marksheet data and builds a structured, beautifully styled Excel workbook.
    Returns bytes of the .xlsx file.
    """
    wb = openpyxl.Workbook()
    # Define styles
    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="FFFFFF")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    regular_font = Font(name="Calibri", size=11)
    bold_font = Font(name="Calibri", size=11, bold=True)
    metric_title_font = Font(name="Calibri", size=9, bold=True, color="595959")
    metric_val_font = Font(name="Calibri", size=12, bold=True, color="1F4E79")

    primary_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    secondary_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    accent_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    zebra_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

    thin_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3")
    )
    double_bottom_border = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="double", color="1F4E79")
    )

    is_broadsheet = data.get("is_broadsheet", False) and bool(data.get("candidates"))

    if is_broadsheet:
        candidates = data.get("candidates", [])
        ws1 = wb.active
        ws1.title = "Cohort Broadsheet"
        ws1.views.sheetView[0].showGridLines = True

        # Header Title Banner
        institution = data.get("institution") or "ACADEMIC MARKSHEET BROADSHEET"
        ws1.merge_cells("A1:J1")
        ws1["A1"] = institution
        ws1["A1"].font = title_font
        ws1["A1"].fill = primary_fill
        ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws1.row_dimensions[1].height = 28

        # Subtitle
        ws1.merge_cells("A2:J2")
        ws1["A2"] = f"Total Candidates Evaluated: {len(candidates)} | Exported from Marksheet Extractor"
        ws1["A2"].font = subtitle_font
        ws1["A2"].fill = secondary_fill
        ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
        ws1.row_dimensions[2].height = 20

        # Collect unique subjects
        subject_names = []
        for cand in candidates:
            for s in cand.get("subjects", []):
                sname = s.get("subject", "Subject")
                if sname not in subject_names:
                    subject_names.append(sname)

        # Table Column Headers
        headers = ["#", "Roll Number", "Student Name"] + subject_names + ["Aggregate", "SGPA", "Result"]
        ws1.append([])  # Row 3 blank
        ws1.append(headers)  # Row 4

        header_row_idx = 4
        ws1.row_dimensions[header_row_idx].height = 26
        for col_idx in range(1, len(headers) + 1):
            cell = ws1.cell(row=header_row_idx, column=col_idx)
            cell.font = header_font
            cell.fill = primary_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        # Populate Student Rows
        for i, cand in enumerate(candidates, start=1):
            roll = get_value(cand.get("roll_no"))
            name = get_value(cand.get("name"))
            agg = get_value(cand.get("aggregate"))
            sgpa = get_value(cand.get("sgpa"))
            res = get_value(cand.get("result"))

            # Build subject map
            cand_subs = {s.get("subject"): s for s in cand.get("subjects", [])}
            sub_values = []
            for sname in subject_names:
                if sname in cand_subs:
                    s_info = cand_subs[sname]
                    obt = s_info.get("obtained", "—")
                    grd = s_info.get("grade")
                    val_str = f"{obt} ({grd})" if grd else f"{obt}"
                    sub_values.append(val_str)
                else:
                    sub_values.append("—")

            row_data = [i, roll, name] + sub_values + [agg, sgpa, res]
            ws1.append(row_data)

            row_idx = header_row_idx + i
            ws1.row_dimensions[row_idx].height = 22
            row_fill = zebra_fill if i % 2 == 0 else PatternFill(fill_type=None)

            for col_idx in range(1, len(headers) + 1):
                cell = ws1.cell(row=row_idx, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if row_fill.fill_type:
                    cell.fill = row_fill

                # Alignment
                if col_idx == 3:  # Name
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    cell.font = bold_font
                elif col_idx in (1, 2):  # Sl no, Roll
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="center", vertical="center")

        # Sheet 2: Detailed Breakdown
        ws2 = wb.create_sheet(title="Subject Details")
        ws2.views.sheetView[0].showGridLines = True
        ws2.merge_cells("A1:G1")
        ws2["A1"] = "INDIVIDUAL SUBJECT-WISE BREAKDOWN"
        ws2["A1"].font = title_font
        ws2["A1"].fill = primary_fill
        ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws2.row_dimensions[1].height = 26

        ws2_headers = ["Roll No", "Student Name", "Subject", "Marks Obtained", "Grade", "Max Marks", "Confidence"]
        ws2.append([])
        ws2.append(ws2_headers)
        ws2.row_dimensions[3].height = 24

        for col_idx in range(1, len(ws2_headers) + 1):
            cell = ws2.cell(row=3, column=col_idx)
            cell.font = header_font
            cell.fill = secondary_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        s2_row = 4
        for cand in candidates:
            roll = get_value(cand.get("roll_no"))
            name = get_value(cand.get("name"))
            for s in cand.get("subjects", []):
                ws2.append([
                    roll,
                    name,
                    s.get("subject", "—"),
                    s.get("obtained", "—"),
                    s.get("grade", "—"),
                    s.get("max", "—"),
                    s.get("confidence", 1.0)
                ])
                for col_idx in range(1, len(ws2_headers) + 1):
                    c = ws2.cell(row=s2_row, column=col_idx)
                    c.font = regular_font
                    c.border = thin_border
                    c.alignment = Alignment(horizontal="left" if col_idx in (2, 3) else "center", vertical="center")
                s2_row += 1

        # Auto-width for Sheet 2
        for col in ws2.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws2.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # Auto-width for Sheet 1
        for col in ws1.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws1.column_dimensions[col_letter].width = max(max_len + 4, 12)

    else:
        # Single Marksheet Extraction
        ws = wb.active
        ws.title = "Student Marksheet"
        ws.views.sheetView[0].showGridLines = True

        name = get_value(data.get("name"))
        roll = get_value(data.get("roll_no"))
        result = get_value(data.get("result"))
        sgpa = get_value(data.get("sgpa", data.get("aggregate", "—")))

        # Title
        ws.merge_cells("A1:E1")
        ws["A1"] = "ACADEMIC MARKSHEET REPORT"
        ws["A1"].font = title_font
        ws["A1"].fill = primary_fill
        ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 30

        # Info Cards Grid
        ws["A3"] = "Candidate Name:"
        ws["A3"].font = bold_font
        ws["B3"] = str(name)
        ws["B3"].font = bold_font

        ws["D3"] = "Roll Number:"
        ws["D3"].font = bold_font
        ws["E3"] = str(roll)
        ws["E3"].font = bold_font

        ws["A4"] = "Overall Result:"
        ws["A4"].font = bold_font
        ws["B4"] = str(result)
        ws["B4"].font = bold_font

        ws["D4"] = "SGPA / Aggregate:"
        ws["D4"].font = bold_font
        ws["E4"] = str(sgpa)
        ws["E4"].font = bold_font

        for r in (3, 4):
            for c in ("A", "B", "D", "E"):
                ws[f"{c}{r}"].fill = accent_fill
                ws[f"{c}{r}"].border = thin_border
            ws.row_dimensions[r].height = 22

        # Subject Table
        headers = ["Sl No", "Subject Name", "Marks Obtained", "Maximum Marks", "Grade / Score"]
        ws.append([])  # Row 5 blank
        ws.append(headers)  # Row 6

        header_row_idx = 6
        ws.row_dimensions[header_row_idx].height = 25
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row_idx, column=col_idx)
            cell.font = header_font
            cell.fill = secondary_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        subjects = data.get("subjects", [])
        total_obtained = 0
        total_max = 0
        has_numeric_totals = False

        for i, s in enumerate(subjects, start=1):
            s_name = s.get("subject", "—")
            s_obt = s.get("obtained", "—")
            s_max = s.get("max", "100")
            s_grd = s.get("grade", "—")

            # Try tracking numeric totals
            try:
                total_obtained += float(s_obt)
                total_max += float(s_max)
                has_numeric_totals = True
            except (ValueError, TypeError):
                pass

            row_data = [i, s_name, s_obt, s_max, s_grd]
            ws.append(row_data)

            curr_row = header_row_idx + i
            ws.row_dimensions[curr_row].height = 20
            row_fill = zebra_fill if i % 2 == 0 else PatternFill(fill_type=None)

            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=curr_row, column=col_idx)
                cell.font = regular_font
                cell.border = thin_border
                if row_fill.fill_type:
                    cell.fill = row_fill
                cell.alignment = Alignment(horizontal="left" if col_idx == 2 else "center", vertical="center")

        # Summary Row
        if has_numeric_totals:
            sum_row_idx = header_row_idx + len(subjects) + 1
            ws.cell(row=sum_row_idx, column=1, value="")
            ws.cell(row=sum_row_idx, column=2, value="TOTAL")
            ws.cell(row=sum_row_idx, column=3, value=round(total_obtained, 2))
            ws.cell(row=sum_row_idx, column=4, value=round(total_max, 2))
            pct = f"{round((total_obtained / total_max) * 100, 2)}%" if total_max else "—"
            ws.cell(row=sum_row_idx, column=5, value=pct)

            ws.row_dimensions[sum_row_idx].height = 24
            for col_idx in range(1, len(headers) + 1):
                c = ws.cell(row=sum_row_idx, column=col_idx)
                c.font = bold_font
                c.fill = accent_fill
                c.border = double_bottom_border
                c.alignment = Alignment(horizontal="left" if col_idx == 2 else "center", vertical="center")

        # Auto-width
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 15)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
